"""Integration tests for the format converters, asserting real outputs."""

from __future__ import annotations

from pathlib import Path

import pytest

from src.engine import pdf_converters
from src.engine.errors import (
    EncryptedPDFError,
    InvalidPasswordError,
    TaskCancelled,
    YourDocError,
)
from tests.support import builders


@pytest.mark.integration
def test_pdf_to_images_png_and_jpeg(scratch: Path, sample_pdf: Path) -> None:
    images = pdf_converters.pdf_to_images(
        str(sample_pdf), scratch / "out" / "images", dpi=72, pages="1-2"
    )
    assert len(images) == 2
    assert all(p.exists() and p.stat().st_size > 0 for p in images)
    assert all(p.suffix == ".png" for p in images)

    jpegs = pdf_converters.pdf_to_images(
        str(sample_pdf), scratch / "out" / "images_jpg", image_format="jpeg", dpi=72, pages="1"
    )
    assert jpegs[0].suffix == ".jpg"


@pytest.mark.integration
def test_pdf_to_word_produces_an_editable_document(scratch: Path, sample_pdf: Path) -> None:
    """The Word output must be a real .docx carrying the PDF's text."""

    out = scratch / "out" / "converted.docx"
    result = pdf_converters.pdf_to_word(str(sample_pdf), out)

    assert result == out
    assert out.exists() and out.stat().st_size > 0

    import docx

    document = docx.Document(str(out))
    body = "\n".join(p.text for p in document.paragraphs)
    assert "CLIENT-Jane Doe" in body, f"text layer missing from the .docx: {body!r}"
    # Four pages were selected, so every page's marker must be present.
    for page in range(1, 5):
        assert f"Page {page} of 4" in body, f"page {page} missing from the .docx"


@pytest.mark.integration
def test_pdf_to_word_honours_the_page_window(scratch: Path, sample_pdf: Path) -> None:
    out = scratch / "out" / "subset.docx"
    pdf_converters.pdf_to_word(str(sample_pdf), out, pages="2-3")

    import docx

    body = "\n".join(p.text for p in docx.Document(str(out)).paragraphs)
    assert "Page 2 of 4" in body
    assert "Page 3 of 4" in body
    assert "Page 1 of 4" not in body, "a page outside the window was converted"
    assert "Page 4 of 4" not in body, "a page outside the window was converted"


@pytest.mark.integration
def test_pdf_to_word_carries_tables_through(scratch: Path, table_pdf: Path) -> None:
    out = scratch / "out" / "tables.docx"
    pdf_converters.pdf_to_word(str(table_pdf), out)

    import docx

    document = docx.Document(str(out))
    assert document.tables, "the detected table did not reach the .docx"
    cells = {cell.text.strip() for row in document.tables[0].rows for cell in row.cells}
    for needle in ("Client", "Acme Ltd", "INV-4471"):
        assert needle in cells, f"{needle!r} missing from the converted table: {cells}"


@pytest.mark.integration
def test_pdf_to_word_refuses_a_pageless_selection(scratch: Path, sample_pdf: Path) -> None:
    out = scratch / "out" / "never.docx"
    with pytest.raises(YourDocError):
        pdf_converters.pdf_to_word(str(sample_pdf), out, pages="99")
    assert not out.exists(), "a document was left behind after failing"


@pytest.mark.integration
def test_pdf_to_word_reports_a_scanned_pdf(scratch: Path) -> None:
    """A PDF with no text layer must be refused, not written as an empty .docx.

    Regression: the converter used to build the document first and return it
    regardless, so a scan produced a valid-looking .docx that opened to a
    blank page and the user had no idea why.
    """

    blank = scratch / "scanned.pdf"
    pdf_converters.images_to_pdf([builders.make_png(scratch / "scan.png")], blank)

    out = scratch / "out" / "scan.docx"
    with pytest.raises(YourDocError) as info:
        pdf_converters.pdf_to_word(str(blank), out)
    message = info.value.user_message.lower()
    assert "scan" in message or "ocr" in message, message
    assert not out.exists(), "an empty .docx was written for a scanned PDF"


@pytest.mark.integration
def test_pdf_to_word_opens_an_encrypted_pdf(scratch: Path, encrypted_pdf: Path) -> None:
    out = scratch / "out" / "locked.docx"
    pdf_converters.pdf_to_word(str(encrypted_pdf), out, password="secret")

    import docx

    body = "\n".join(p.text for p in docx.Document(str(out)).paragraphs)
    assert "CLIENT-Jane Doe" in body


@pytest.mark.integration
def test_pdf_to_word_asks_for_the_password_on_a_locked_file(
    scratch: Path, encrypted_pdf: Path
) -> None:
    """A locked PDF must raise the specific error, not a generic failure.

    Regression: pdfplumber surfaces a missing password as pdfminer's
    ``PDFPasswordIncorrect``, which is not part of our hierarchy -- so the
    converter's blanket handler rewrapped it as ``OperationFailedError``
    ("may be encrypted, or may have an unsupported layout"). The user was left
    to infer that the password field was the answer, while Merge and PDF to
    Images already told them outright.
    """

    out = scratch / "out" / "never.docx"

    with pytest.raises(EncryptedPDFError) as info:
        pdf_converters.pdf_to_word(str(encrypted_pdf), out)
    assert "password" in info.value.user_message.lower()
    assert not out.exists(), "a document was written for a locked input"

    with pytest.raises(InvalidPasswordError):
        pdf_converters.pdf_to_word(str(encrypted_pdf), out, password="wrong")


@pytest.mark.integration
def test_pdf_to_excel_asks_for_the_password_on_a_locked_file(
    scratch: Path, encrypted_table_pdf: Path
) -> None:
    """The same, for the Excel converter."""

    out = scratch / "out" / "never.xlsx"

    with pytest.raises(EncryptedPDFError) as info:
        pdf_converters.pdf_to_excel(str(encrypted_table_pdf), out)
    assert "password" in info.value.user_message.lower()
    assert not out.exists(), "a workbook was written for a locked input"

    with pytest.raises(InvalidPasswordError):
        pdf_converters.pdf_to_excel(
            str(encrypted_table_pdf), out, password="wrong"
        )


@pytest.mark.integration
def test_pdf_to_word_writes_atomically(scratch: Path, sample_pdf: Path) -> None:
    out = scratch / "out"
    pdf_converters.pdf_to_word(str(sample_pdf), out / "clean.docx")
    pdf_converters.pdf_to_word(str(sample_pdf), out / "clean.docx")
    assert not _temp_files(out), f"temp files leaked: {_temp_files(out)}"


@pytest.mark.integration
def test_pdf_to_word_cancel_leaves_nothing_behind(scratch: Path, sample_pdf: Path) -> None:
    import threading

    cancel = threading.Event()
    seen: list[float] = []

    def progress(fraction: float, _message: str) -> None:
        seen.append(fraction)
        if len(seen) >= 2:
            cancel.set()

    out = scratch / "out" / "cancelled.docx"
    with pytest.raises((TaskCancelled, YourDocError)):
        pdf_converters.pdf_to_word(
            str(sample_pdf), out, pages="1-4", progress=progress, cancel_event=cancel
        )
    assert not out.exists(), "a cancelled run left a document behind"


@pytest.mark.integration
def test_images_to_pdf_fit_and_a4(scratch: Path, sample_pdf: Path) -> None:
    images = pdf_converters.pdf_to_images(str(sample_pdf), scratch / "out" / "imgs", dpi=72, pages="1-3")

    rebuilt = pdf_converters.images_to_pdf(images, scratch / "out" / "rebuilt.pdf")
    assert builders.page_count(rebuilt) == 3

    a4 = pdf_converters.images_to_pdf(images, scratch / "out" / "rebuilt_a4.pdf", page_size="a4")
    assert builders.page_count(a4) == 3


@pytest.mark.integration
def test_images_to_pdf_rejects_empty(scratch: Path) -> None:
    with pytest.raises(YourDocError):
        pdf_converters.images_to_pdf([], scratch / "out" / "empty.pdf")


@pytest.mark.integration
def test_images_to_pdf_honours_order(scratch: Path) -> None:
    red = builders.make_png(scratch / "red.png", color=(255, 0, 0))
    blue = builders.make_png(scratch / "blue.png", color=(0, 0, 255))
    out = pdf_converters.images_to_pdf([red, blue], scratch / "out" / "ordered.pdf")
    assert builders.page_count(out) == 2
    # Each page must carry its own image's colour, which is what makes this a
    # test of order rather than just of page count.
    first = builders.page_mean_rgb(out, page=0)
    second = builders.page_mean_rgb(out, page=1)
    assert first[0] > first[2], f"page 1 is not red-dominant: {first}"
    assert second[2] > second[0], f"page 2 is not blue-dominant: {second}"


@pytest.mark.integration
def test_pdf_to_excel_with_tables(scratch: Path, table_pdf: Path) -> None:
    xlsx = scratch / "out" / "tables.xlsx"
    _, table_count = pdf_converters.pdf_to_excel(str(table_pdf), xlsx)
    assert table_count >= 1

    import openpyxl

    workbook = openpyxl.load_workbook(str(xlsx))
    text = " | ".join(
        str(cell.value)
        for sheet in workbook.worksheets
        for row in sheet.iter_rows()
        for cell in row
        if cell.value is not None
    )
    for needle in ("Client", "Acme Ltd", "INV-4471", "Beta Corp", "980.00"):
        assert needle in text, f"{needle!r} missing from workbook: {text}"


@pytest.mark.integration
def test_pdf_to_excel_no_tables_refused(scratch: Path, sample_pdf: Path) -> None:
    target = scratch / "out" / "empty.xlsx"
    with pytest.raises(YourDocError) as info:
        pdf_converters.pdf_to_excel(str(sample_pdf), target)
    assert "table" in info.value.user_message.lower()
    assert not target.exists(), "a workbook was left behind after failing"


# --------------------------------------------------------------------------- #
# Atomic output
# --------------------------------------------------------------------------- #
def _temp_files(directory: Path) -> list[Path]:
    """Every leftover temp file in ``directory``."""

    return list(directory.glob("*.yourdoc.tmp*"))


@pytest.mark.integration
def test_converters_leave_no_temp_files(scratch: Path, sample_pdf: Path, table_pdf: Path) -> None:
    """Regression: the PDF converters wrote straight to the destination.

    Only the Office converters used the atomic helper, so a cancelled or failed
    run could leave a half-written file where the user expected a whole one.
    Each converter now writes to a temp sibling and promotes it on success.
    """

    out = scratch / "out"
    pdf_converters.pdf_to_excel(str(table_pdf), out / "a.xlsx")
    pdf_converters.pdf_to_excel(str(table_pdf), out / "b.xlsx")
    pdf_converters.images_to_pdf([builders.make_png(scratch / "p.png")], out / "c.pdf")
    pdf_converters.pdf_to_images(str(sample_pdf), out / "imgs", dpi=72, pages="1")

    assert not _temp_files(out), f"temp files leaked: {_temp_files(out)}"
    assert not _temp_files(out / "imgs"), "temp files leaked in the image folder"


@pytest.mark.integration
def test_pdf_to_excel_replaces_an_existing_file_atomically(
    scratch: Path, table_pdf: Path
) -> None:
    """A rerun replaces the previous workbook and leaves no debris."""

    target = scratch / "out" / "rerun.xlsx"
    pdf_converters.pdf_to_excel(str(table_pdf), target)
    first_size = target.stat().st_size
    pdf_converters.pdf_to_excel(str(table_pdf), target)
    assert target.exists()
    assert target.stat().st_size == first_size
    assert not _temp_files(target.parent)


@pytest.mark.integration
def test_pdf_to_images_rolls_back_pages_on_cancel(scratch: Path, sample_pdf: Path) -> None:
    """Cancelling mid-render must not leave a half-exported set of pages."""

    import threading

    out = scratch / "out" / "imgs"
    cancel = threading.Event()

    # Cancel once the first page has been written: the progress callback fires
    # before each render, so letting the first through and stopping the second
    # exercises the rollback path.
    seen: list[float] = []

    def progress(fraction: float, _message: str) -> None:
        seen.append(fraction)
        if len(seen) >= 2:
            cancel.set()

    with pytest.raises((TaskCancelled, YourDocError)):
        pdf_converters.pdf_to_images(
            str(sample_pdf), out, dpi=72, pages="1-4", progress=progress, cancel_event=cancel
        )

    leftovers = sorted(out.glob("*.png")) if out.exists() else []
    assert not leftovers, f"cancelled export left pages behind: {leftovers}"
    if out.exists():
        assert not _temp_files(out), f"temp files leaked: {_temp_files(out)}"


@pytest.mark.integration
@pytest.mark.parametrize("dpi", [0, 35, 5000])
def test_pdf_to_images_rejects_out_of_range_dpi(
    scratch: Path, sample_pdf: Path, dpi: int
) -> None:
    """A caller-supplied DPI must not drive an unbounded raster allocation."""

    with pytest.raises(YourDocError):
        pdf_converters.pdf_to_images(
            str(sample_pdf), scratch / "out" / "imgs", dpi=dpi, pages="1"
        )
