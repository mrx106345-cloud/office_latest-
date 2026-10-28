"""Caps that turn a pathological input into a clear message instead of a hang.

Each one bounds something that was previously unbounded, so the app either
refuses early or says plainly that the output is incomplete. A cap that
silently truncated would be worse than no cap: the result is a valid, normal
looking document.
"""

from __future__ import annotations

from pathlib import Path

import pytest

from src.engine import office, pdf_converters
from src.engine.errors import InvalidInputError, YourDocError
from tests.support import builders


# --------------------------------------------------------------------------- #
# B1 - worksheet XML part size
# --------------------------------------------------------------------------- #
@pytest.mark.unit
def test_oversized_sheet_xml_is_skipped_without_being_parsed(
    scratch: Path, monkeypatch
) -> None:
    """Regression: the whole sheet part was read to reach a few KB of ``<cols>``.

    Measured at ~2.8 KB per row for a 50-column sheet, so a 100,000-row sheet is
    ~276 MB materialised and then expanded into millions of Element objects.
    ``MemoryError`` was not in the handled tuple, so it escaped as a generic
    "This is unexpected" -- or took the process down.
    """

    import zipfile

    source = builders.make_xlsx(scratch / "w.xlsx")

    # A normal sheet keeps its real column widths.
    widths = office._workbook_column_widths(source, [source.stem])
    assert widths, "the fixture should produce column widths"

    # Drop the cap to nothing: the part is refused on its declared size, so the
    # part is never decompressed, and the caller falls back to Excel defaults.
    original = office._MAX_SHEET_XML_BYTES
    monkeypatch.setattr(office, "_MAX_SHEET_XML_BYTES", 1)
    try:
        assert office._workbook_column_widths(source, [source.stem]) == {
            source.stem: {}
        }
    finally:
        monkeypatch.setattr(office, "_MAX_SHEET_XML_BYTES", original)

    # And it still works once the cap is back.
    assert office._workbook_column_widths(source, [source.stem])


@pytest.mark.unit
def test_sheet_xml_cap_is_far_above_real_workloads() -> None:
    """The cap must not reject work a user would reasonably expect to convert."""

    # 10,000 rows at ~2.8 KB/row is the largest legitimate sheet admitted.
    assert office._MAX_SHEET_XML_BYTES >= 10_000 * 2762
    # ...and it is still a real ceiling, not "effectively unlimited".
    assert office._MAX_SHEET_XML_BYTES <= 512 * 1024 * 1024


# --------------------------------------------------------------------------- #
# B2 - rows and columns
# --------------------------------------------------------------------------- #
@pytest.mark.integration
def test_large_sheet_is_truncated_and_says_so(scratch: Path, monkeypatch) -> None:
    """Regression: an unbounded sheet exhausted memory or froze for minutes.

    The conversion still happens -- the user's request is honoured as far as is
    sensible -- but the result is marked as incomplete, in the file itself and
    in the returned note.
    """

    import openpyxl

    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet.title = "Big"
    for index in range(40):
        sheet.append([f"r{index}c{column}" for column in range(4)])
    source = scratch / "big.xlsx"
    workbook.save(source)

    monkeypatch.setattr(office, "_MAX_SHEET_ROWS", 10)

    notes: list = []
    result = office.excel_to_pdf(source, scratch / "out.pdf", notes=notes)

    assert result.exists()
    assert len(notes) == 1, "truncation was not reported"
    note = notes[0]
    assert note.truncated
    assert note.rows_total == 40
    assert note.rows_shown == 10
    message = note.describe()
    assert "10" in message and "40" in message
    assert "incomplete" in message.lower()
    # The sheet is named, so a multi-sheet workbook is unambiguous.
    assert "Big" in message


@pytest.mark.integration
def test_truncation_note_is_written_into_the_pdf(scratch: Path, monkeypatch) -> None:
    """The file leaves the app and looks normal in any reader, so it must say so."""

    import openpyxl

    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet.title = "Big"
    for index in range(40):
        sheet.append([f"r{index}c{column}" for column in range(4)])
    source = scratch / "big.xlsx"
    workbook.save(source)

    monkeypatch.setattr(office, "_MAX_SHEET_ROWS", 10)

    result = office.excel_to_pdf(source, scratch / "out.pdf", notes=[])
    text = builders.extract_all_text(result)
    assert "Note:" in text, "the PDF does not record that it is incomplete"


@pytest.mark.integration
def test_truncation_note_is_written_into_the_docx(scratch: Path, monkeypatch) -> None:
    import docx
    import openpyxl

    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet.title = "Big"
    for index in range(40):
        sheet.append([f"r{index}c{column}" for column in range(4)])
    source = scratch / "big.xlsx"
    workbook.save(source)

    monkeypatch.setattr(office, "_MAX_SHEET_ROWS", 10)

    result = office.excel_to_word(source, scratch / "out.docx", notes=[])
    document = docx.Document(str(result))
    texts = [p.text for p in document.paragraphs if "Note:" in p.text]
    assert texts, "the .docx does not record that it is incomplete"


@pytest.mark.unit
def test_a_sheet_within_the_caps_reports_nothing(scratch: Path) -> None:
    """No note for ordinary input -- the warning has to mean something."""

    import openpyxl

    workbook = openpyxl.Workbook()
    workbook.active["A1"] = "a"
    source = scratch / "small.xlsx"
    workbook.save(source)

    notes: list = []
    office._used_rows(
        openpyxl.load_workbook(source, read_only=True).active, notes=notes
    )
    assert notes == [], f"a small sheet was reported as truncated: {notes}"


@pytest.mark.unit
def test_wide_sheet_is_truncated_on_columns_too(scratch: Path) -> None:
    """A row cap alone does not bound memory: iter_rows pads to max_column."""

    import openpyxl

    workbook = openpyxl.Workbook()
    sheet = workbook.active
    # Wider than the cap. 39 columns would sit comfortably inside it.
    for column in range(1, office._MAX_SHEET_COLUMNS + 40):
        sheet.cell(row=1, column=column, value=f"c{column}")
    source = scratch / "wide.xlsx"
    workbook.save(source)

    notes: list = []
    rows = office._used_rows(
        openpyxl.load_workbook(source, read_only=True).active, notes=notes
    )
    assert len(notes) == 1, f"a sheet over the column cap was not reported: {notes}"
    assert notes[0].columns_total > office._MAX_SHEET_COLUMNS
    assert notes[0].columns_shown == office._MAX_SHEET_COLUMNS
    assert all(len(row) == office._MAX_SHEET_COLUMNS for row in rows)
    assert "columns" in notes[0].describe()


# --------------------------------------------------------------------------- #
# B3 - total decoded pixels
# --------------------------------------------------------------------------- #
@pytest.mark.integration
def test_too_many_pixels_is_refused_before_decoding(scratch: Path, monkeypatch) -> None:
    """Regression: every image was decoded and held until the final save.

    Pillow guards a single image, but nothing guarded the aggregate: a decoded
    12-megapixel photo costs ~35 MB, so 100 of them is ~3.5 GB before any
    conversion overhead.
    """

    images = builders.make_images(6, scratch / "imgs")
    # Each fixture is 200x140, so this budget admits a handful and no more.
    monkeypatch.setattr(pdf_converters, "_MAX_TOTAL_PIXELS", 6 * 200 * 140 - 1)

    with pytest.raises(InvalidInputError) as info:
        pdf_converters.images_to_pdf(images, scratch / "out.pdf")
    assert "too many or too large" in info.value.user_message.lower()
    assert not (scratch / "out.pdf").exists(), "a refused batch wrote an output"


@pytest.mark.integration
def test_image_batch_within_budget_still_converts(scratch: Path) -> None:
    images = builders.make_images(6, scratch / "imgs")
    result = pdf_converters.images_to_pdf(images, scratch / "out.pdf")
    assert builders.page_count(result) == 6


@pytest.mark.unit
def test_pixel_cap_admits_the_existing_fixture_with_margin() -> None:
    """The perf suite builds 60 images; the cap must sit well above that."""

    fixture_pixels = 60 * 200 * 140
    assert pdf_converters._MAX_TOTAL_PIXELS > fixture_pixels * 20


# --------------------------------------------------------------------------- #
# Stage A - the one unwrapped write
# --------------------------------------------------------------------------- #
@pytest.mark.integration
def test_pdf_to_images_reports_a_bad_destination(scratch: Path, sample_pdf: Path) -> None:
    """Regression: the mkdir ran before any handler and leaked a raw OSError.

    Every other converter reported a clean message for the same destination;
    this one raised a bare FileExistsError, which the UI shows as "This is
    unexpected".
    """

    blocker = scratch / "not_a_folder.txt"
    blocker.write_text("x", encoding="utf-8")

    with pytest.raises(YourDocError) as info:
        pdf_converters.pdf_to_images(
            str(sample_pdf), str(blocker / "images"), dpi=36, pages="1"
        )
    assert "writable" in info.value.user_message.lower() or "folder" in info.value.user_message.lower()