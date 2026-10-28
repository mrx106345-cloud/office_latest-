"""Integration tests for the four office converters, asserting real outputs."""

from __future__ import annotations

import threading
from pathlib import Path

import pytest

from src.engine import office
from src.engine.errors import InvalidInputError, OperationFailedError, TaskCancelled
from tests.support import builders

def _record(calls: list[tuple[float, str]]):
    """A progress callback that captures every ``(fraction, message)`` pair."""

    def report(fraction: float, message: str) -> None:
        calls.append((fraction, message))

    return report

# --------------------------------------------------------------------------- #
# Word -> PDF
# --------------------------------------------------------------------------- #
@pytest.mark.integration
def test_word_to_pdf_renders_every_paragraph_style(scratch: Path) -> None:
    source = builders.make_docx(scratch / "report.docx", page_break=True)
    out = scratch / "out" / "report.pdf"

    result = office.word_to_pdf(source, out)

    assert result == out
    assert result.stat().st_size > 0
    assert builders.page_count(result) >= 2  # the page break started a new one
    text = builders.extract_all_text(result)
    for needle in (
        "Annual report",
        "Heading one",
        "Heading two",
        "Heading three",
        "Regular",
        "bold",
        "italic",
        "sized",
        "First bullet",
        "Second bullet",
        "Nested bullet",
        "First number",
        "Second number",
        "Numbered level two",
        "A quoted phrase",
        "An intense quote",
        "Figure 1",
        "Centred line",
        "Right-aligned line",
        "After a page break",
        "Client",
        "Amount",
        "Reference",
        "Acme Ltd 1",
        "100.50",
        "INV-0001",
    ):
        assert needle in text, f"missing from the PDF: {needle}"
    assert not list(out.parent.glob("*.yourdoc.tmp.*"))

@pytest.mark.integration
def test_word_to_pdf_repeats_the_table_header_on_every_page(scratch: Path) -> None:
    source = builders.make_docx(scratch / "tall.docx", tall_table_rows=120)
    out = scratch / "out" / "tall.pdf"

    result = office.word_to_pdf(source, out)

    assert builders.page_count(result) > 2
    text = builders.extract_all_text(result)
    assert text.count("Client") > 1  # the shaded header row repeats
    assert "Acme Ltd 120" in text
    assert "INV-0120" in text

@pytest.mark.integration
def test_word_to_pdf_reports_progress(scratch: Path) -> None:
    source = builders.make_docx(scratch / "report.docx")
    calls: list[tuple[float, str]] = []

    office.word_to_pdf(
        source, scratch / "out" / "p.pdf", progress=_record(calls)
    )

    fractions = [fraction for fraction, _ in calls]
    assert fractions[0] == pytest.approx(0.02)
    assert fractions == sorted(fractions)
    assert fractions[-1] == pytest.approx(1.0)
    assert any("Rendering" in message for _, message in calls)
    assert "Saved" in calls[-1][1]

@pytest.mark.integration
def test_word_to_pdf_rejects_an_empty_document(scratch: Path) -> None:
    source = builders.make_empty_docx(scratch / "empty.docx")
    out = scratch / "out" / "empty.pdf"

    with pytest.raises(OperationFailedError) as info:
        office.word_to_pdf(source, out)

    assert "empty" in info.value.user_message.lower()
    assert "empty.docx" in info.value.user_message
    assert not out.exists()
    assert not list(out.parent.glob("*.yourdoc.tmp.*"))

@pytest.mark.integration
def test_word_to_pdf_rejects_a_corrupt_file(scratch: Path) -> None:
    source = builders.make_text_file(
        scratch / "corrupt.docx", content="PK\x03\x04garbage"
    )
    out = scratch / "out" / "corrupt.pdf"

    with pytest.raises(OperationFailedError) as info:
        office.word_to_pdf(source, out)

    assert "corrupt" in info.value.user_message.lower()
    assert not out.exists()

@pytest.mark.integration
def test_word_to_pdf_honours_cancellation(scratch: Path) -> None:
    source = builders.make_docx(scratch / "report.docx")
    out = scratch / "out" / "cancelled.pdf"
    event = threading.Event()
    event.set()

    with pytest.raises(TaskCancelled):
        office.word_to_pdf(source, out, cancel_event=event)

    assert not out.exists()
    assert not list(out.parent.glob("*.yourdoc.tmp.*"))

@pytest.mark.integration
def test_word_to_pdf_validates_its_input(scratch: Path, sample_pdf: Path) -> None:
    out = scratch / "out" / "checked.pdf"
    with pytest.raises(InvalidInputError):
        office.word_to_pdf(scratch / "ghost.docx", out)
    with pytest.raises(InvalidInputError):
        office.word_to_pdf(sample_pdf, out)  # a PDF is not a Word file
    assert not out.exists()

# --------------------------------------------------------------------------- #
# Excel -> PDF
# --------------------------------------------------------------------------- #
@pytest.mark.integration
def test_excel_to_pdf_renders_every_sheet_in_landscape(scratch: Path) -> None:
    source = builders.make_xlsx(scratch / "book.xlsx")
    out = scratch / "out" / "book.pdf"

    result = office.excel_to_pdf(source, out)

    assert result == out
    assert result.stat().st_size > 0
    sizes = builders.page_sizes(result)
    assert len(sizes) == 2  # one landscape page per sheet
    for width, height in sizes:
        assert width > height, f"page is not landscape: {width}x{height}"
    text = builders.extract_all_text(result)
    for needle in (
        "Region",
        "North",
        "South",
        "East",
        "42",
        "1200.5",
        "89.25",
        "TRUE",
        "FALSE",
        "2024-05-01T10:30:00",
        "2024-05-02",
    ):
        assert needle in text, f"missing from the PDF: {needle}"
    assert not list(out.parent.glob("*.yourdoc.tmp.*"))

@pytest.mark.integration
def test_excel_to_pdf_scales_a_wide_sheet_down(scratch: Path) -> None:
    source = builders.make_xlsx(scratch / "wide.xlsx", wide=True)
    out = scratch / "out" / "wide.pdf"

    result = office.excel_to_pdf(source, out)

    assert builders.page_count(result) == 3  # two narrow sheets + the wide one
    text = builders.extract_all_text(result)
    assert "Region" in text
    assert "Column" in text and "value" in text  # scaled cells still render

@pytest.mark.integration
def test_excel_to_pdf_reports_progress(scratch: Path) -> None:
    source = builders.make_xlsx(scratch / "book.xlsx")
    calls: list[tuple[float, str]] = []

    office.excel_to_pdf(
        source, scratch / "out" / "p.pdf", progress=_record(calls)
    )

    fractions = [fraction for fraction, _ in calls]
    assert fractions[0] == pytest.approx(0.02)
    assert fractions == sorted(fractions)
    assert fractions[-1] == pytest.approx(1.0)
    assert any("Rendering sheet" in message for _, message in calls)

@pytest.mark.integration
def test_excel_to_pdf_rejects_a_workbook_without_data(scratch: Path) -> None:
    source = builders.make_empty_xlsx(scratch / "empty.xlsx")
    out = scratch / "out" / "empty.pdf"

    with pytest.raises(OperationFailedError) as info:
        office.excel_to_pdf(source, out)

    assert "no data" in info.value.user_message.lower()
    assert not out.exists()
    assert not list(out.parent.glob("*.yourdoc.tmp.*"))

@pytest.mark.integration
def test_excel_to_pdf_rejects_a_corrupt_file(scratch: Path) -> None:
    source = builders.make_text_file(scratch / "corrupt.xlsx", content="not a zip")
    out = scratch / "out" / "corrupt.pdf"

    with pytest.raises(OperationFailedError) as info:
        office.excel_to_pdf(source, out)

    assert "corrupt" in info.value.user_message.lower()
    assert not out.exists()

@pytest.mark.integration
def test_excel_to_pdf_honours_cancellation(scratch: Path) -> None:
    source = builders.make_xlsx(scratch / "book.xlsx")
    out = scratch / "out" / "cancelled.pdf"
    event = threading.Event()
    event.set()

    with pytest.raises(TaskCancelled):
        office.excel_to_pdf(source, out, cancel_event=event)

    assert not out.exists()

@pytest.mark.integration
def test_excel_to_pdf_validates_its_input(scratch: Path, sample_pdf: Path) -> None:
    out = scratch / "out" / "checked.pdf"
    with pytest.raises(InvalidInputError):
        office.excel_to_pdf(scratch / "ghost.xlsx", out)
    with pytest.raises(InvalidInputError):
        office.excel_to_pdf(sample_pdf, out)  # a PDF is not an Excel file
    assert not out.exists()

# --------------------------------------------------------------------------- #
# Word -> Excel
# --------------------------------------------------------------------------- #
@pytest.mark.integration
def test_word_to_excel_maps_text_and_tables_to_sheets(scratch: Path) -> None:
    source = builders.make_docx(scratch / "report.docx")
    out = scratch / "out" / "report.xlsx"

    result = office.word_to_excel(source, out)

    assert result == out
    assert result.stat().st_size > 0

    import openpyxl

    workbook = openpyxl.load_workbook(str(result))
    assert workbook.sheetnames == ["Text", "Table 1"]

    text_sheet = workbook["Text"]
    cells = [row[0].value for row in text_sheet.iter_rows()]
    title = cells.index("Annual report")
    heading = cells.index("Heading one")
    body = next(
        cell
        for cell in cells
        if isinstance(cell, str) and cell.startswith("Regular text")
    )
    assert "bold run" in body and "italic run" in body
    assert text_sheet.cell(title + 1, 1).font.bold
    assert text_sheet.cell(heading + 1, 1).font.bold
    assert text_sheet.column_dimensions["A"].width == 100

    table_sheet = workbook["Table 1"]
    assert table_sheet["A1"].value == "Client"
    assert table_sheet["A2"].value == "Acme Ltd 1"
    assert table_sheet["B2"].value == "100.50"
    assert table_sheet["C4"].value == "INV-0003"
    assert table_sheet["A1"].font.bold
    assert table_sheet["A1"].fill.patternType == "solid"
    assert table_sheet.column_dimensions["A"].width == 12
    assert not list(out.parent.glob("*.yourdoc.tmp.*"))

@pytest.mark.integration
def test_word_to_excel_from_a_text_only_document(scratch: Path) -> None:
    import docx

    source = scratch / "notes.docx"
    document = docx.Document()
    document.add_paragraph("Just a note")
    document.add_heading("A heading", level=2)
    document.save(str(source))
    out = scratch / "out" / "notes.xlsx"

    result = office.word_to_excel(source, out)

    import openpyxl

    workbook = openpyxl.load_workbook(str(result))
    assert workbook.sheetnames == ["Text"]
    assert workbook["Text"]["A1"].value == "Just a note"
    assert workbook["Text"]["A2"].value == "A heading"
    assert workbook["Text"]["A2"].font.bold

@pytest.mark.integration
def test_word_to_excel_reports_progress(scratch: Path) -> None:
    source = builders.make_docx(scratch / "report.docx")
    calls: list[tuple[float, str]] = []

    office.word_to_excel(
        source, scratch / "out" / "p.xlsx", progress=_record(calls)
    )

    fractions = [fraction for fraction, _ in calls]
    assert fractions == sorted(fractions)
    assert fractions[-1] == pytest.approx(1.0)
    assert any("Writing table" in message for _, message in calls)

@pytest.mark.integration
def test_word_to_excel_rejects_an_empty_document(scratch: Path) -> None:
    source = builders.make_empty_docx(scratch / "empty.docx")
    out = scratch / "out" / "empty.xlsx"

    with pytest.raises(OperationFailedError) as info:
        office.word_to_excel(source, out)

    assert "no text or tables" in info.value.user_message.lower()
    assert not out.exists()

@pytest.mark.integration
def test_word_to_excel_rejects_a_corrupt_file(scratch: Path) -> None:
    source = builders.make_text_file(
        scratch / "corrupt.docx", content="PK\x03\x04garbage"
    )
    out = scratch / "out" / "corrupt.xlsx"

    with pytest.raises(OperationFailedError) as info:
        office.word_to_excel(source, out)

    assert "corrupt" in info.value.user_message.lower()
    assert not out.exists()

@pytest.mark.integration
def test_word_to_excel_honours_cancellation(scratch: Path) -> None:
    source = builders.make_docx(scratch / "report.docx")
    out = scratch / "out" / "cancelled.xlsx"
    event = threading.Event()
    event.set()

    with pytest.raises(TaskCancelled):
        office.word_to_excel(source, out, cancel_event=event)

    assert not out.exists()

@pytest.mark.integration
def test_word_to_excel_validates_its_input(scratch: Path, sample_pdf: Path) -> None:
    out = scratch / "out" / "checked.xlsx"
    with pytest.raises(InvalidInputError):
        office.word_to_excel(scratch / "ghost.docx", out)
    with pytest.raises(InvalidInputError):
        office.word_to_excel(sample_pdf, out)  # a PDF is not a Word file
    assert not out.exists()

# --------------------------------------------------------------------------- #
# Excel -> Word
# --------------------------------------------------------------------------- #
@pytest.mark.integration
def test_excel_to_word_maps_sheets_to_headed_tables(scratch: Path) -> None:
    source = builders.make_xlsx(scratch / "book.xlsx")
    out = scratch / "out" / "book.docx"

    result = office.excel_to_word(source, out)

    assert result == out
    assert result.stat().st_size > 0

    import docx

    document = docx.Document(str(result))
    headings = [
        paragraph.text
        for paragraph in document.paragraphs
        if paragraph.style.name.startswith("Heading")
    ]
    assert headings == ["Sheet 1", "Sheet 2"]
    assert len(document.tables) == 2

    table = document.tables[0]
    assert table.cell(0, 0).text == "Region"
    assert table.cell(1, 0).text == "North"
    assert table.cell(1, 3).text == "TRUE"
    assert table.cell(2, 4).text == "2024-05-02"
    assert all(
        run.bold
        for cell in table.rows[0].cells
        for paragraph in cell.paragraphs
        for run in paragraph.runs
    )
    assert not list(out.parent.glob("*.yourdoc.tmp.*"))

@pytest.mark.integration
def test_excel_to_word_chunks_large_sheets(scratch: Path) -> None:
    source = builders.make_xlsx(
        scratch / "big.xlsx", big_rows=office._CHUNK_ROWS + 10
    )
    out = scratch / "out" / "big.docx"

    result = office.excel_to_word(source, out)

    import docx

    document = docx.Document(str(result))
    headings = [
        paragraph.text
        for paragraph in document.paragraphs
        if paragraph.style.name.startswith("Heading")
    ]
    assert headings[-2:] == ["Big", "Big (continued)"]
    assert len(document.tables) == 4  # two plain sheets + the chunked one
    first, second = document.tables[-2:]
    assert len(first.rows) == office._CHUNK_ROWS
    assert len(second.rows) == 11  # the header plus the remaining rows

@pytest.mark.integration
def test_excel_to_word_reports_progress(scratch: Path) -> None:
    source = builders.make_xlsx(scratch / "book.xlsx")
    calls: list[tuple[float, str]] = []

    office.excel_to_word(
        source, scratch / "out" / "p.docx", progress=_record(calls)
    )

    fractions = [fraction for fraction, _ in calls]
    assert fractions == sorted(fractions), "progress went backwards"
    assert fractions[-1] == pytest.approx(1.0)
    # Progress now names the sheet being written, row by row, rather than once
    # per sheet. That granularity is the fix: a big sheet used to sit at a
    # frozen bar for minutes with an inert Cancel button.
    messages = [message for _, message in calls]
    assert any("Writing" in message for message in messages), messages
    assert any("Reading sheet" in message for message in messages), messages

@pytest.mark.integration
def test_excel_to_word_rejects_a_workbook_without_data(scratch: Path) -> None:
    source = builders.make_empty_xlsx(scratch / "empty.xlsx")
    out = scratch / "out" / "empty.docx"

    with pytest.raises(OperationFailedError) as info:
        office.excel_to_word(source, out)

    assert "no data" in info.value.user_message.lower()
    assert not out.exists()

@pytest.mark.integration
def test_excel_to_word_rejects_a_corrupt_file(scratch: Path) -> None:
    source = builders.make_text_file(scratch / "corrupt.xlsx", content="not a zip")
    out = scratch / "out" / "corrupt.docx"

    with pytest.raises(OperationFailedError) as info:
        office.excel_to_word(source, out)

    assert "corrupt" in info.value.user_message.lower()
    assert not out.exists()

@pytest.mark.integration
def test_excel_to_word_honours_cancellation(scratch: Path) -> None:
    source = builders.make_xlsx(scratch / "book.xlsx")
    out = scratch / "out" / "cancelled.docx"
    event = threading.Event()
    event.set()

    with pytest.raises(TaskCancelled):
        office.excel_to_word(source, out, cancel_event=event)

    assert not out.exists()

@pytest.mark.integration
def test_excel_to_word_validates_its_input(scratch: Path, sample_pdf: Path) -> None:
    out = scratch / "out" / "checked.docx"
    with pytest.raises(InvalidInputError):
        office.excel_to_word(scratch / "ghost.xlsx", out)
    with pytest.raises(InvalidInputError):
        office.excel_to_word(sample_pdf, out)  # a PDF is not an Excel file
    assert not out.exists()

# --------------------------------------------------------------------------- #
# Round trip
# --------------------------------------------------------------------------- #
@pytest.mark.integration
def test_word_excel_word_round_trip(scratch: Path) -> None:
    source = builders.make_docx(scratch / "report.docx")
    workbook_path = scratch / "out" / "report.xlsx"
    office.word_to_excel(source, workbook_path)
    back = scratch / "out" / "roundtrip.docx"

    result = office.excel_to_word(workbook_path, back)

    import docx

    document = docx.Document(str(result))
    headings = [
        paragraph.text
        for paragraph in document.paragraphs
        if paragraph.style.name.startswith("Heading")
    ]
    assert headings == ["Text", "Table 1"]
    table = document.tables[1]
    assert table.cell(0, 0).text == "Client"
    assert table.cell(1, 0).text == "Acme Ltd 1"
    assert table.cell(3, 2).text == "INV-0003"


def test_covered_scripts_survive_and_missing_ones_are_reported(scratch):
    """Cyrillic and Greek are drawn; scripts the fallback lacks are reported.

    The base-14 fonts cover Latin-1 only. DejaVu Sans, the bundled fallback,
    adds Cyrillic and Greek but carries no CJK glyphs, and drawing a missing
    glyph silently drops it -- so the export must name what it left out.
    Expecting CJK to be drawn is how this suite used to break after the move
    off MuPDF.
    """

    import docx
    import pypdf

    samples = ["\u041e\u0422\u0427\u0451\u0422 2026", "\u0388\u03ba\u03b8\u03b5\u03c3\u03b7", "\u62a5\u544a\u6c47\u603b"]
    document = docx.Document()
    for sample in samples:
        document.add_paragraph(sample)
    source = scratch / "unicode.docx"
    document.save(str(source))

    result = office.word_to_pdf(source, scratch / "unicode.pdf")

    text = "\n".join(
        page.extract_text() or "" for page in pypdf.PdfReader(str(result)).pages
    )

    # What the fallback covers reaches the page, extractable, not placeholder dots.
    assert "\u041e\u0422\u0427\u0451\u0422" in text, text
    assert "2026" in text
    assert "\u0388\u03ba\u03b8\u03b5\u03c3\u03b7" in text, text

    # And what it cannot draw is *named*, not silently dropped. Code points rather
    # than the glyphs, because the glyphs are what cannot be shown.
    assert "could not be drawn" in text, text
    assert "U+62A5" in text, text


def test_covered_scripts_are_inked_not_just_encoded(scratch):
    """Cyrillic and Greek must be *drawn*, not merely present in the content stream.

    ``test_covered_scripts_survive_and_missing_ones_are_reported`` extracts the
    text and finds the characters there, which is necessary but not sufficient:
    a glyph the font cannot draw is still emitted as a text-showing operator, so
    the extraction passes while the page renders blank. Only rendering the page
    and looking for ink distinguishes the two.

    The document holds nothing but the non-Latin lines, so any ink at all has to
    have come from them -- no band arithmetic to get wrong, and no Latin text to
    be mistaken for a success.
    """

    import docx

    document = docx.Document()
    document.add_paragraph("ОТЧЁТ 2026")
    document.add_paragraph("Έκθεση")
    source = scratch / "inked.docx"
    document.save(str(source))

    result = office.word_to_pdf(source, scratch / "inked.pdf")

    darkest = builders.darkest_ink(result)
    assert darkest < 100, (
        f"the page is blank (darkest pixel {darkest}/255): the Cyrillic and "
        "Greek characters were encoded into the content stream but no glyphs "
        "were drawn for them"
    )


def test_latin_only_pdf_embeds_no_fallback_font(scratch):
    """A Latin-only document embeds no Unicode font at all.

    DejaVu is a third of a megabyte; paying for it on every document that never
    needs it would inflate every export the tool produces.
    """

    import docx
    import pypdf

    document = docx.Document()
    document.add_paragraph("Plain ASCII only, 12345")
    source = scratch / "latin.docx"
    document.save(str(source))

    result = office.word_to_pdf(source, scratch / "latin.pdf")

    embedded: set[str] = set()
    for page in pypdf.PdfReader(str(result)).pages:
        resources = page["/Resources"]  # resolves the indirect reference
        fonts = resources.get("/Font")
        if hasattr(fonts, "get_object"):  # the font dict is usually indirect too
            fonts = fonts.get_object()
        names = list(fonts.values()) if isinstance(fonts, dict) else []
        for reference in names:
            try:
                obj = reference.get_object()
                base = obj.get("/BaseFont", "")
            except Exception:  # noqa: BLE001 - one broken ref must not fail the scan
                continue
            embedded.add(str(base).split("+")[-1])

    assert not any("DejaVu" in name for name in embedded), embedded
    assert embedded, "expected at least one font reference"
def test_latin_only_pdf_still_uses_the_compact_base_fonts(scratch):
    """The Unicode font is embedded on demand, not on every page."""

    import docx

    document = docx.Document()
    document.add_paragraph("Plain ASCII only, 12345")
    source = scratch / "latin.docx"
    document.save(str(source))

    result = office.word_to_pdf(source, scratch / "latin.pdf")

    names = builders.page_font_names(result)
    assert names <= {"Helvetica", "Helvetica-Bold"}, names


def test_a_row_taller_than_a_page_is_split_not_clipped(scratch):
    """An over-long table row must paginate instead of overflowing the page.

    A single cell can wrap into more lines than fit on one page. Drawing it in
    one piece put content below the bottom margin, where it was silently
    clipped -- measured 853pt of content on an 842pt page.
    """

    import docx

    words = "alpha beta gamma delta epsilon zeta eta theta iota kappa".split()
    document = docx.Document()
    table = document.add_table(rows=2, cols=1)
    # Each row is long enough to need its own page, so the repeated header has
    # to be accounted for separately.
    table.cell(0, 0).text = "H " + " ".join(words * 60)
    table.cell(1, 0).text = "D " + " ".join(words * 60)
    source = scratch / "tall.docx"
    document.save(str(source))

    result = office.word_to_pdf(source, scratch / "tall.pdf")

    total_pages = builders.page_count(result)
    assert total_pages > 1, "expected the tall table to span pages"
    sizes = builders.page_sizes(result)

    # The table borders are the thing that could spill: a cell drawn past the
    # bottom edge is silently clipped rather than paginated. A page with no
    # vector operators at all (a plain text page) is skipped, but the run is
    # asserted non-empty so this check cannot pass by inspecting nothing.
    checked = 0
    for index in range(total_pages):
        extents = builders.page_drawn_extents(result, index)
        if extents is None:
            continue
        checked += 1
        assert extents[1] <= sizes[index][1] + 1, (
            f"page {index + 1} drew to y={extents[1]:.1f} on a "
            f"{sizes[index][1]:.1f}pt page"
        )
    assert checked, "no vector operators found - the overflow check was vacuous"

    extracted = builders.extract_all_text(result)
    for word in set(words):
        assert word in extracted, f"{word!r} was lost while splitting the row"

