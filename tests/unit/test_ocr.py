"""M2 OCR: Tesseract wrapper and the scanned-PDF -> Word path.

The native ``tesseract`` binary is a buyer install, so the engine logic is
unit-tested with ``pytesseract`` monkeypatched (mirroring how the optional
dependency is surfaced) and the real-binary cases are skipped when it is absent.
The integration test drives the full :func:`pdf_to_word` OCR branch with the
recognizer stubbed to fixed text, so it exercises detection + wiring without a
Tesseract install.
"""

from __future__ import annotations

import shutil
from pathlib import Path
from types import SimpleNamespace

import pytest

from src.engine import ocr, pdf_converters
from src.engine.errors import DependencyMissingError
from tests.support import builders as b


# --------------------------------------------------------------------------- #
# ocr.py wrapper (pytesseract stubbed -- no native binary needed)
# --------------------------------------------------------------------------- #
def test_ocr_image_calls_the_wrapper_with_lang(monkeypatch: pytest.MonkeyPatch) -> None:
    calls: dict = {}

    def fake_image_to_string(image, lang="eng"):  # noqa: ARG001
        calls["lang"] = lang
        return "recognized text\n"

    fake = SimpleNamespace(
        image_to_string=fake_image_to_string,
        get_tesseract_version=lambda: "5.4.0",
    )
    monkeypatch.setattr(ocr, "require_pytesseract", lambda: fake)

    assert ocr.ocr_image(object(), lang="spa").strip() == "recognized text"
    assert calls["lang"] == "spa"


def test_tesseract_version_surfaces_missing_binary(monkeypatch: pytest.MonkeyPatch) -> None:
    def boom():
        raise RuntimeError("TesseractNotFoundError")

    fake = SimpleNamespace(get_tesseract_version=boom)
    monkeypatch.setattr(ocr, "require_pytesseract", lambda: fake)

    with pytest.raises(DependencyMissingError) as info:
        ocr.tesseract_version()
    # The friendly message carries per-OS install guidance, not a raw crash.
    assert "Tesseract" in info.value.user_message


def test_tesseract_available_is_false_when_version_raises(monkeypatch: pytest.MonkeyPatch) -> None:
    monkeypatch.setattr(ocr, "tesseract_version", lambda: (_ for _ in ()).throw(DependencyMissingError("x")))
    assert ocr.tesseract_available() is False


def test_tesseract_available_is_true_when_version_ok(monkeypatch: pytest.MonkeyPatch) -> None:
    monkeypatch.setattr(ocr, "tesseract_version", lambda: "5.4.0")
    assert ocr.tesseract_available() is True


# --------------------------------------------------------------------------- #
# pdf_to_word(ocr=True): detection + wiring, recognizer stubbed
# --------------------------------------------------------------------------- #
def _image_only_pdf(tmp_path: Path) -> Path:
    """A PDF built from a picture, so it carries no text layer to read."""

    scan = tmp_path / "scan.pdf"
    b.make_png(tmp_path / "scan.png")
    pdf_converters.images_to_pdf([tmp_path / "scan.png"], scan)
    return scan


def test_pdf_to_word_ocr_recognizes_a_scanned_page(tmp_path: Path, monkeypatch: pytest.MonkeyPatch) -> None:
    source = _image_only_pdf(tmp_path)
    out = tmp_path / "ocr.docx"

    monkeypatch.setattr(ocr, "render_page_to_image", lambda doc, idx, dpi=300: object())
    monkeypatch.setattr(ocr, "ocr_image", lambda image, lang="eng": "SCANNED-CONTRACT-4471")

    result = pdf_converters.pdf_to_word(source, out, ocr=True)
    assert result == out and out.exists()

    import docx

    body = "\n".join(p.text for p in docx.Document(str(out)).paragraphs)
    assert "SCANNED-CONTRACT-4471" in body, body


def test_pdf_to_word_ocr_skips_pages_with_a_text_layer(tmp_path: Path, monkeypatch: pytest.MonkeyPatch) -> None:
    """A page that already has text must never be rasterized for OCR."""

    source = b.make_sample_pdf(tmp_path / "doc.pdf", pages=2)
    out = tmp_path / "word.docx"

    def _never(*args, **kwargs):  # noqa: ARG001
        raise AssertionError("OCR was attempted on a text-layer page")

    monkeypatch.setattr(pdf_converters, "_ocr_page_text", _never)

    pdf_converters.pdf_to_word(source, out, ocr=True)

    import docx

    body = "\n".join(p.text for p in docx.Document(str(out)).paragraphs)
    assert "CLIENT-Jane Doe" in body


def test_pdf_to_word_without_ocr_refuses_a_scan(tmp_path: Path) -> None:
    """The default (non-OCR) path still refuses an empty scan, pointing at OCR."""

    source = _image_only_pdf(tmp_path)
    out = tmp_path / "nope.docx"

    with pytest.raises(Exception) as info:
        pdf_converters.pdf_to_word(source, out)
    assert not out.exists()
    assert "ocr" in str(getattr(info.value, "user_message", info.value)).lower()


# --------------------------------------------------------------------------- #
# View wiring (no Tk: the language-code helper and registration are pure)
# --------------------------------------------------------------------------- #
def test_ocr_view_is_registered() -> None:
    from src.ui import views

    keys = {view.key for view in views.VIEW_CLASSES}
    assert "pdf_ocr_to_word" in keys


def test_ocr_view_maps_display_to_language_code() -> None:
    from src.ui.views import PdfOcrToWordView

    assert PdfOcrToWordView._lang_code("English (eng)") == "eng"
    assert PdfOcrToWordView._lang_code("Spanish (spa)") == "spa"
    assert PdfOcrToWordView._lang_code("  French  ") == "french"


# --------------------------------------------------------------------------- #
# Real Tesseract, skipped when the binary is absent
# --------------------------------------------------------------------------- #
@pytest.mark.skipif(not shutil.which("tesseract"), reason="tesseract binary not installed")
def test_real_tesseract_reads_a_fixture(tmp_path: Path) -> None:  # pragma: no cover - env dependent
    from PIL import Image, ImageDraw

    image = Image.new("RGB", (400, 120), "white")
    ImageDraw.Draw(image).text((20, 40), "HELLO OCR", fill="black")

    text = ocr.ocr_image(image, lang="eng")
    assert "HELLO" in text.upper()
