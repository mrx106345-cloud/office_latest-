"""Pure helpers behind the four office converters, tested in isolation."""

from __future__ import annotations

import datetime
import threading
from pathlib import Path

import pytest

from src.engine import office
from src.engine.errors import InvalidInputError, OperationFailedError, TaskCancelled
from tests.support import builders


# --------------------------------------------------------------------------- #
# _validate_input
# --------------------------------------------------------------------------- #
@pytest.mark.unit
def test_validate_input_accepts_a_real_word_file(scratch: Path) -> None:
    source = builders.make_docx(scratch / "report.docx")
    assert office._validate_input(source, office.WORD_SUFFIXES, "Word") == source


@pytest.mark.unit
def test_validate_input_accepts_str_and_uppercase_suffix(scratch: Path) -> None:
    source = builders.make_docx(scratch / "REPORT.DOCX")
    resolved = office._validate_input(str(source), office.WORD_SUFFIXES, "Word")
    assert resolved == source


@pytest.mark.unit
def test_validate_input_rejects_missing_file(scratch: Path) -> None:
    ghost = scratch / "ghost.docx"
    with pytest.raises(InvalidInputError) as info:
        office._validate_input(ghost, office.WORD_SUFFIXES, "Word")
    assert "Missing input file" in str(info.value)
    assert ghost.name in info.value.user_message


@pytest.mark.unit
def test_validate_input_rejects_directory(scratch: Path) -> None:
    folder = scratch / "folder.xlsx"
    folder.mkdir()
    with pytest.raises(InvalidInputError) as info:
        office._validate_input(folder, office.EXCEL_SUFFIXES, "Excel")
    assert "Not a regular file" in str(info.value)
    assert "folder" in info.value.user_message.lower()


@pytest.mark.unit
def test_validate_input_rejects_wrong_extension(scratch: Path) -> None:
    text = builders.make_text_file(scratch / "notes.txt")
    with pytest.raises(InvalidInputError) as info:
        office._validate_input(text, office.WORD_SUFFIXES, "Word")
    assert "Not a Word file" in str(info.value)
    assert ".docx" in info.value.user_message


@pytest.mark.unit
def test_validate_input_rejects_pdf_as_excel(sample_pdf: Path) -> None:
    with pytest.raises(InvalidInputError):
        office._validate_input(sample_pdf, office.EXCEL_SUFFIXES, "Excel")


# --------------------------------------------------------------------------- #
# _progress_sink
# --------------------------------------------------------------------------- #
@pytest.mark.unit
def test_progress_sink_forwards_to_the_callback() -> None:
    calls: list[tuple[float, str]] = []

    def record(fraction: float, message: str) -> None:
        calls.append((fraction, message))

    sink = office._progress_sink(record, None)
    sink(0.25, "Working...")
    assert calls == [(0.25, "Working...")]


@pytest.mark.unit
def test_progress_sink_is_a_noop_without_a_callback() -> None:
    sink = office._progress_sink(None, None)
    sink(0.5, "Quiet")  # must not raise


@pytest.mark.unit
def test_progress_sink_raises_task_cancelled_on_cancel() -> None:
    """Regression: this used to raise ``OperationFailedError``.

    The task runner routed anything that was not ``TaskCancelled`` to
    ``on_error``, so a cancelled Office run painted a red "Failed" banner
    instead of the neutral cancelled state. Cancellation now uses the same
    exception type as the PDF engines.
    """

    event = threading.Event()
    event.set()
    sink = office._progress_sink(None, event)
    with pytest.raises(TaskCancelled):
        sink(0.1, "Opening document...")


@pytest.mark.unit
def test_progress_sink_cancels_at_the_next_checkpoint() -> None:
    event = threading.Event()
    sink = office._progress_sink(lambda f, m: event.set(), event)
    sink(0.1, "first checkpoint passes")
    with pytest.raises(TaskCancelled):
        sink(0.2, "second checkpoint is cancelled")


@pytest.mark.unit
def test_progress_sink_swallows_callback_errors() -> None:
    def broken(fraction: float, message: str) -> None:
        raise ValueError("broken sink")

    sink = office._progress_sink(broken, None)
    sink(0.5, "the broken sink must not break the run")


# --------------------------------------------------------------------------- #
# _stringify / _local_name
# --------------------------------------------------------------------------- #
@pytest.mark.unit
@pytest.mark.parametrize(
    ("value", "expected"),
    [
        (None, ""),
        (True, "TRUE"),
        (False, "FALSE"),
        (42, "42"),
        (-7, "-7"),
        (3.14, "3.14"),
        (1e-7, "1e-07"),
        (0.1 + 0.2, "0.3"),
        (datetime.datetime(2024, 5, 1, 10, 30), "2024-05-01T10:30:00"),
        (datetime.date(2024, 5, 1), "2024-05-01"),
        (datetime.time(10, 30, 5), "10:30:05"),
        ("text", "text"),
        (object(), None),  # anything else falls back to str()
    ],
)
def test_stringify_renders_cell_values(value: object, expected: str | None) -> None:
    rendered = office._stringify(value)
    if expected is None:
        assert rendered == str(value)
    else:
        assert rendered == expected


@pytest.mark.unit
def test_local_name_strips_xml_namespaces() -> None:
    assert office._local_name("{http://example.com}col") == "col"
    assert office._local_name("col") == "col"
    assert office._local_name("{}p") == "p"


# --------------------------------------------------------------------------- #
# _TextMetrics
# --------------------------------------------------------------------------- #
@pytest.fixture
def metrics():
    return office._TextMetrics()


@pytest.mark.unit
def test_metrics_width_matches_reportlab(metrics: office._TextMetrics) -> None:
    from reportlab.pdfbase import pdfmetrics

    assert metrics.width("hello", "helv", 10) == pdfmetrics.stringWidth(
        "hello", "Helvetica", 10
    )
    assert metrics.width("hello", "helv", 10) > 0
    assert metrics.width("hello", "hebo", 10) > metrics.width("hello", "helv", 10)


@pytest.mark.unit
def test_metrics_line_width_counts_spaces(metrics: office._TextMetrics) -> None:
    line = [("hello", "helv", 10), ("world", "helv", 10)]
    expected = (
        metrics.width("hello", "helv", 10)
        + metrics.width(" ", "helv", 10)
        + metrics.width("world", "helv", 10)
    )
    assert metrics.line_width(line) == pytest.approx(expected)
    assert metrics.line_width(line) > metrics.width("helloworld", "helv", 10)


@pytest.mark.unit
def test_metrics_wrap_single_line(metrics: office._TextMetrics) -> None:
    lines = metrics.wrap([("hello world", "helv", 10)], 500.0)
    assert len(lines) == 1
    assert [token[0] for token in lines[0]] == ["hello", "world"]


@pytest.mark.unit
def test_metrics_wrap_breaks_on_newlines(metrics: office._TextMetrics) -> None:
    lines = metrics.wrap([("one\ntwo", "helv", 10)], 500.0)
    assert len(lines) == 2
    assert [token[0] for token in lines[0]] == ["one"]
    assert [token[0] for token in lines[1]] == ["two"]


@pytest.mark.unit
def test_metrics_wrap_drops_trailing_empty_line(metrics: office._TextMetrics) -> None:
    assert len(metrics.wrap([("hello\n", "helv", 10)], 500.0)) == 1


@pytest.mark.unit
def test_metrics_wrap_of_nothing_is_one_empty_line(
    metrics: office._TextMetrics,
) -> None:
    assert metrics.wrap([], 500.0) == [[]]


@pytest.mark.unit
def test_metrics_wrap_hard_splits_overwide_words(
    metrics: office._TextMetrics,
) -> None:
    word = "supercalifragilisticexpialidocious"
    limit = 30.0
    lines = metrics.wrap([(word, "helv", 10)], limit)
    assert len(lines) > 1
    joined = "".join(token[0] for line in lines for token in line)
    assert joined == word  # no characters lost
    for line in lines:
        assert metrics.line_width(line) <= limit


@pytest.mark.unit
def test_metrics_wrap_mixes_segments_and_sizes(
    metrics: office._TextMetrics,
) -> None:
    segments = [("bold ", "hebo", 12.0), ("plain", "helv", 10.0)]
    lines = metrics.wrap(segments, 500.0)
    assert len(lines) == 1
    assert lines[0] == [
        ("bold", "hebo", 12.0),
        ("plain", "helv", 10.0),
    ]


@pytest.mark.unit
def test_metrics_split_returns_the_word_untouched_when_fitting(
    metrics: office._TextMetrics,
) -> None:
    assert metrics._split("hi", "helv", 10, 500.0) == ["hi"]


@pytest.mark.unit
def test_metrics_split_chunks_a_giant_word(metrics: office._TextMetrics) -> None:
    chunks = metrics._split("abcdefghij" * 10, "helv", 10, 20.0)
    assert len(chunks) > 1
    assert all(metrics.width(chunk, "helv", 10) <= 20.0 for chunk in chunks)
    assert "".join(chunks) == "abcdefghij" * 10


# --------------------------------------------------------------------------- #
# _distribute
# --------------------------------------------------------------------------- #
@pytest.mark.unit
def test_distribute_splits_evenly_when_nothing_is_known() -> None:
    assert office._distribute([None, None, None], 90.0) == [30.0, 30.0, 30.0]
    assert office._distribute([None], 50.0) == [50.0]


@pytest.mark.unit
def test_distribute_gives_leftover_to_unknown_columns() -> None:
    assert office._distribute([60.0, None], 100.0) == [60.0, 40.0]
    assert office._distribute([20.0, None, None], 80.0) == [20.0, 30.0, 30.0]


@pytest.mark.unit
def test_distribute_keeps_known_columns_when_they_fit() -> None:
    assert office._distribute([20.0, 30.0], 100.0) == [20.0, 30.0]


@pytest.mark.unit
def test_distribute_scales_an_overwide_grid_down() -> None:
    assert office._distribute([80.0, 80.0], 100.0) == [50.0, 50.0]
    assert office._distribute([60.0, None], 40.0) == [40.0, 0.0]


# --------------------------------------------------------------------------- #
# _sheet_name
# --------------------------------------------------------------------------- #
@pytest.mark.unit
def test_sheet_name_keeps_valid_names() -> None:
    assert office._sheet_name("Sales", set()) == "Sales"
    assert office._sheet_name("Sheet 1", set()) == "Sheet 1"


@pytest.mark.unit
def test_sheet_name_caps_at_thirty_one_characters() -> None:
    long_name = "x" * 50
    assert office._sheet_name(long_name, set()) == "x" * 31


@pytest.mark.unit
def test_sheet_name_replaces_invalid_characters() -> None:
    cleaned = office._sheet_name("Q1/2024 *report?", set())
    assert cleaned == "Q1_2024 _report_"
    assert len(cleaned) <= 31


@pytest.mark.unit
def test_sheet_name_falls_back_to_sheet_for_blanks() -> None:
    assert office._sheet_name("", set()) == "Sheet"
    assert office._sheet_name("   ", set()) == "Sheet"


@pytest.mark.unit
def test_sheet_name_dedupes_case_insensitively() -> None:
    taken: set[str] = set()
    assert office._sheet_name("Data", taken) == "Data"
    assert office._sheet_name("data", taken) == "data_1"
    assert office._sheet_name("DATA", taken) == "DATA_2"


@pytest.mark.unit
def test_sheet_name_keeps_deduplicated_names_within_the_cap() -> None:
    taken = {"x" * 31}
    deduped = office._sheet_name("x" * 31, taken)
    assert len(deduped) == 31
    assert deduped.endswith("_1")


# --------------------------------------------------------------------------- #
# _excel_width_to_points
# --------------------------------------------------------------------------- #
@pytest.mark.unit
def test_excel_width_to_points_uses_the_digit_width_formula() -> None:
    # 8.43 characters -> 59 px + 5 px padding -> 64 px -> 48 pt
    assert office._excel_width_to_points(8.43) == pytest.approx(48.0)
    assert office._excel_width_to_points(1.0) == pytest.approx(9.0)
    assert office._excel_width_to_points(2.5) == pytest.approx(17.25)


# --------------------------------------------------------------------------- #
# _used_rows
# --------------------------------------------------------------------------- #
@pytest.mark.unit
def test_used_rows_trims_trailing_empty_rows_and_columns() -> None:
    import openpyxl

    worksheet = openpyxl.Workbook().active
    worksheet["A1"] = 1
    worksheet["A2"] = None  # explicit empty cell: extends the used range
    worksheet["C1"] = None  # a column that is empty everywhere
    assert office._used_rows(worksheet) == [[1]]


@pytest.mark.unit
def test_used_rows_keeps_interior_blank_rows() -> None:
    import openpyxl

    worksheet = openpyxl.Workbook().active
    worksheet["A1"] = 1
    worksheet["A3"] = 3
    assert office._used_rows(worksheet) == [[1], [None], [3]]


@pytest.mark.unit
def test_used_rows_trims_appended_empty_rows() -> None:
    import openpyxl

    worksheet = openpyxl.Workbook().active
    worksheet.append([1, 2])
    worksheet.append([None, None])
    worksheet.append([None, None])
    assert office._used_rows(worksheet) == [[1, 2]]


@pytest.mark.unit
def test_used_rows_of_an_empty_sheet_is_empty() -> None:
    import openpyxl

    assert office._used_rows(openpyxl.Workbook().active) == []


# --------------------------------------------------------------------------- #
# _workbook_column_widths / _sheet_xml_targets
# --------------------------------------------------------------------------- #
@pytest.mark.unit
def test_sheet_xml_targets_maps_every_sheet(scratch: Path) -> None:
    import openpyxl
    import zipfile

    workbook = openpyxl.Workbook()
    workbook.active.title = "First"
    workbook.create_sheet("Second")
    path = scratch / "targets.xlsx"
    workbook.save(str(path))

    with zipfile.ZipFile(str(path)) as archive:
        targets = office._sheet_xml_targets(archive)

    assert set(targets) == {"First", "Second"}
    assert targets["First"].endswith("sheet1.xml")
    assert targets["Second"].endswith("sheet2.xml")


@pytest.mark.unit
def test_sheet_xml_targets_of_a_package_without_the_workbook_part(
    scratch: Path,
) -> None:
    import zipfile

    path = scratch / "no_workbook.zip"
    with zipfile.ZipFile(path, "w") as archive:
        archive.writestr("readme.txt", "not an xlsx package")

    with zipfile.ZipFile(path) as archive:
        assert office._sheet_xml_targets(archive) == {}


@pytest.mark.unit
def test_sheet_xml_targets_skips_sheets_without_a_name(scratch: Path) -> None:
    import zipfile

    path = scratch / "unnamed.xlsx"
    _craft_workbook(path, sheet_xml="", sheet_name="")

    with zipfile.ZipFile(path) as archive:
        assert office._sheet_xml_targets(archive) == {}


@pytest.mark.unit
def test_workbook_column_widths_reads_the_cols_elements(scratch: Path) -> None:
    import openpyxl

    workbook = openpyxl.Workbook()
    first = workbook.active
    first.title = "First"
    first.append(["a", "b"])
    first.column_dimensions["A"].width = 22
    first.column_dimensions["B"].width = 9.5
    second = workbook.create_sheet("Second")
    second.append(["x"])  # no explicit widths anywhere
    path = scratch / "widths.xlsx"
    workbook.save(str(path))

    widths = office._workbook_column_widths(
        path, ["First", "Second", "Absent"]
    )
    assert widths["First"] == {1: 22.0, 2: 9.5}
    assert widths["Second"] == {}
    assert widths["Absent"] == {}


@pytest.mark.unit
def test_workbook_column_widths_ignores_unrequested_sheets(scratch: Path) -> None:
    import openpyxl

    workbook = openpyxl.Workbook()
    workbook.active.title = "First"
    workbook.create_sheet("Second")
    path = scratch / "subset.xlsx"
    workbook.save(str(path))

    assert office._workbook_column_widths(path, ["Second"]) == {"Second": {}}


def _craft_workbook(
    path: Path,
    *,
    sheet_xml: str,
    include_sheet_part: bool = True,
    sheet_name: str = "Data",
) -> None:
    """Write a minimal one-sheet package straight to ``path``."""

    import zipfile

    workbook = (
        '<?xml version="1.0"?>'
        f'<workbook xmlns="{office._SS_NS}" xmlns:r="{office._REL_NS}">'
        f'<sheets><sheet name="{sheet_name}" sheetId="1" r:id="rId1"/></sheets>'
        "</workbook>"
    )
    relationships = (
        '<?xml version="1.0"?>'
        f'<Relationships xmlns="{office._REL_NS}">'
        '<Relationship Id="rId1" '
        f'Type="{office._REL_NS}/worksheet" '
        'Target="worksheets/sheet1.xml"/>'
        "</Relationships>"
    )
    with zipfile.ZipFile(path, "w") as archive:
        archive.writestr("xl/workbook.xml", workbook)
        archive.writestr("xl/_rels/workbook.xml.rels", relationships)
        if include_sheet_part:
            archive.writestr("xl/worksheets/sheet1.xml", sheet_xml)


@pytest.mark.unit
def test_workbook_column_widths_skips_broken_column_elements(scratch: Path) -> None:
    sheet = (
        '<?xml version="1.0"?>'
        f'<worksheet xmlns="{office._SS_NS}"'
        "><cols>"
        '<col min="abc" max="2" width="9"/>'  # min is not a number
        '<col min="0" max="2" width="9"/>'  # min must be positive
        '<col min="3" max="2" width="9"/>'  # max before min
        '<col min="4" max="4" width="-3"/>'  # negative width
        '<col min="5" max="6" width="12.5"/>'  # valid
        "</cols><sheetData/></worksheet>"
    )
    path = scratch / "crafted.xlsx"
    _craft_workbook(path, sheet_xml=sheet)

    assert office._workbook_column_widths(path, ["Data"]) == {
        "Data": {5: 12.5, 6: 12.5}
    }


@pytest.mark.unit
def test_workbook_column_widths_skips_a_missing_sheet_part(scratch: Path) -> None:
    path = scratch / "missing_part.xlsx"
    _craft_workbook(path, sheet_xml="", include_sheet_part=False)

    assert office._workbook_column_widths(path, ["Data"]) == {"Data": {}}


@pytest.mark.unit
def test_workbook_column_widths_of_a_non_package(scratch: Path) -> None:
    junk = builders.make_text_file(scratch / "junk.xlsx")

    assert office._workbook_column_widths(junk, ["Data"]) == {"Data": {}}


# --------------------------------------------------------------------------- #
# _autosize_columns
# --------------------------------------------------------------------------- #
@pytest.mark.unit
def test_autosize_columns_sizes_from_the_longest_value() -> None:
    import openpyxl

    worksheet = openpyxl.Workbook().active
    office._autosize_columns(
        worksheet,
        [["Name", "Notes"], ["Alice", "x" * 100], ["Bob", None]],
        openpyxl,
    )
    assert worksheet.column_dimensions["A"].width == 8  # "Alice" + 2, clamped
    assert worksheet.column_dimensions["B"].width == 60  # clamped at the top


@pytest.mark.unit
def test_autosize_columns_ignores_rows_shorter_than_the_widest() -> None:
    import openpyxl

    worksheet = openpyxl.Workbook().active
    office._autosize_columns(worksheet, [["a", "b", "ccc"], ["x"]], openpyxl)
    assert worksheet.column_dimensions["C"].width == 8  # "ccc" + 2, clamped


@pytest.mark.unit
def test_autosize_columns_leaves_empty_input_alone() -> None:
    import openpyxl

    worksheet = openpyxl.Workbook().active
    office._autosize_columns(worksheet, [], openpyxl)
    assert "A" not in worksheet.column_dimensions


# --------------------------------------------------------------------------- #
# _add_sheet_to_docx
# --------------------------------------------------------------------------- #
@pytest.mark.unit
def test_add_sheet_to_docx_writes_one_hearded_table() -> None:
    import docx

    document = docx.Document()
    rows = [["Region", "Units"], ["North", "42"], ["South", "7"]]
    office._add_sheet_to_docx(document, "Data", rows)

    assert [para.text for para in document.paragraphs] == ["Data"]
    assert len(document.tables) == 1
    table = document.tables[0]
    assert table.cell(0, 0).text == "Region"
    assert table.cell(1, 1).text == "42"
    header_runs = [
        run
        for cell in table.rows[0].cells
        for paragraph in cell.paragraphs
        for run in paragraph.runs
    ]
    assert header_runs and all(run.bold for run in header_runs)


@pytest.mark.unit
def test_add_sheet_to_docx_chunks_large_sheets() -> None:
    import docx

    document = docx.Document()
    rows = [["Index"]] + [[str(index)] for index in range(office._CHUNK_ROWS + 1)]
    office._add_sheet_to_docx(document, "Big", rows)

    assert [para.text for para in document.paragraphs] == [
        "Big",
        "Big (continued)",
    ]
    assert len(document.tables) == 2
    assert len(document.tables[0].rows) == office._CHUNK_ROWS
    assert len(document.tables[1].rows) == 2


@pytest.mark.unit
def test_add_sheet_to_docx_pads_short_rows() -> None:
    import docx

    document = docx.Document()
    office._add_sheet_to_docx(document, "Ragged", [["A", "B", "C"], ["only"]])
    table = document.tables[0]
    assert table.cell(1, 0).text == "only"
    assert table.cell(1, 1).text == ""
    assert table.cell(1, 2).text == ""


# --------------------------------------------------------------------------- #
# Style-name parsing helpers (stubbed paragraphs)
# --------------------------------------------------------------------------- #
@pytest.mark.unit
def test_style_name_returns_the_style_or_empty() -> None:
    class _Broken:
        @property
        def style(self):
            raise KeyError("dangling styleId")

    assert office._style_name(_Broken()) == ""


@pytest.mark.unit
def test_style_name_handles_a_missing_style() -> None:
    class _Missing:
        style = None

    assert office._style_name(_Missing()) == ""


@pytest.mark.unit
def test_has_numbering_reads_direct_list_formatting() -> None:
    class _PPr:
        numPr = object()

    class _With:
        class _p:
            pPr = _PPr()

    class _Without:
        class _p:
            pPr = None

    assert office._has_numbering(_With()) is True
    assert office._has_numbering(_Without()) is False


@pytest.mark.unit
@pytest.mark.parametrize(
    ("style", "expected"),
    [
        ("heading 1", 1),
        ("heading 2", 2),
        ("heading 3", 3),
        ("heading 9", 3),  # clamped
        ("heading", 1),  # no number
        ("heading x", 1),  # not a number
        ("heading 0", 1),  # clamped from below
    ],
)
def test_heading_level_parses_and_clamps(style: str, expected: int) -> None:
    assert office._heading_level(style) == expected


@pytest.mark.unit
@pytest.mark.parametrize(
    ("style", "expected"),
    [
        ("list bullet 1", 1),
        ("list bullet 2", 2),
        ("list bullet 3", 3),
        ("list bullet 9", 3),  # clamped
        ("list bullet", 1),
        ("list number 2", 2),
        ("body", 1),
    ],
)
def test_list_level_parses_nesting(style: str, expected: int) -> None:
    assert office._list_level(style) == expected


@pytest.mark.unit
def test_alignment_maps_paragraph_alignment() -> None:
    class _Aligned:
        def __init__(self, name: str | None) -> None:
            self.alignment = type("A", (), {"name": name})()

    class _Raising:
        @property
        def alignment(self):
            raise RuntimeError("exotic alignment")

    assert office._alignment(_Aligned("CENTER")) == "center"
    assert office._alignment(_Aligned("RIGHT")) == "right"
    assert office._alignment(_Aligned("LEFT")) == "left"
    assert office._alignment(_Aligned("JUSTIFY")) == "left"
    assert office._alignment(_Aligned(None)) == "left"
    assert office._alignment(_Raising()) == "left"


# --------------------------------------------------------------------------- #
# _rich_segments / _column_width_pt
# --------------------------------------------------------------------------- #
@pytest.mark.unit
def test_rich_segments_flattens_run_formatting() -> None:
    import docx
    from docx.shared import Pt

    paragraph = docx.Document().add_paragraph()
    paragraph.add_run("")  # empty runs are dropped
    paragraph.add_run("plain")
    paragraph.add_run("bold").bold = True
    paragraph.add_run("italic").italic = True
    both = paragraph.add_run("both")
    both.bold = True
    both.italic = True
    paragraph.add_run("sized").font.size = Pt(14)
    paragraph.add_run("with\ta tab")

    segments = office._rich_segments(paragraph, 10.5)
    assert segments == [
        ("plain", "helv", 10.5),
        ("bold", "hebo", 10.5),
        ("italic", "heit", 10.5),
        ("both", "hebi", 10.5),
        ("sized", "helv", 14.0),
        ("with a tab", "helv", 10.5),
    ]


@pytest.mark.unit
def test_rich_segments_of_an_empty_paragraph() -> None:
    import docx

    paragraph = docx.Document().add_paragraph()
    assert office._rich_segments(paragraph, 10.5) == [("", "helv", 10.5)]


@pytest.mark.unit
def test_rich_segments_honours_the_defaults() -> None:
    import docx

    paragraph = docx.Document().add_paragraph()
    paragraph.add_run("bold italic").bold = True
    paragraph.add_run("more").italic = True
    segments = office._rich_segments(paragraph, 12.0, bold=True, italic=True)
    assert segments == [
        ("bold italic", "hebi", 12.0),
        ("more", "hebi", 12.0),
    ]


@pytest.mark.unit
def test_rich_segments_falls_back_when_the_size_is_broken() -> None:
    class _BrokenFont:
        @property
        def size(self) -> None:
            raise TypeError("no size on this run")

    class _StubRun:
        text = "kept"
        bold = False
        italic = False
        font = _BrokenFont()

    class _StubParagraph:
        runs = [_StubRun()]

    assert office._rich_segments(_StubParagraph(), 10.5) == [
        ("kept", "helv", 10.5)
    ]


@pytest.mark.unit
def test_column_width_pt_reads_points_or_none() -> None:
    import docx
    from docx.oxml.ns import qn
    from docx.shared import Inches

    document = docx.Document()
    table = document.add_table(rows=1, cols=2, style="Table Grid")
    table.columns[0].width = Inches(1.5)
    # A grid column without a ``w:w`` attribute has no explicit width.
    del table.columns[1]._gridCol.attrib[qn("w:w")]

    assert office._column_width_pt(table.columns[0]) == pytest.approx(108.0)
    assert office._column_width_pt(table.columns[1]) is None


@pytest.mark.unit
def test_column_width_pt_swallows_broken_widths() -> None:
    class _Broken:
        width = "not a length"

    class _Raising:
        @property
        def width(self):
            raise AttributeError("no width")

    assert office._column_width_pt(_Broken()) is None
    assert office._column_width_pt(_Raising()) is None


# --------------------------------------------------------------------------- #
# _DocRenderer table guards
# --------------------------------------------------------------------------- #
@pytest.mark.unit
def test_renderer_table_skips_a_grid_without_columns() -> None:
    canvas = office._PdfCanvas(*office._A4_PORTRAIT, office._MARGIN)
    if True:
        renderer = office._DocRenderer(canvas)

        class _NoColumns:
            columns: list = []
            rows: list = []

        renderer.table(_NoColumns())  # nothing to draw, no error


@pytest.mark.unit
def test_renderer_table_skips_a_grid_without_rows() -> None:
    canvas = office._PdfCanvas(*office._A4_PORTRAIT, office._MARGIN)
    if True:
        renderer = office._DocRenderer(canvas)

        class _NoRows:
            columns = [object(), object()]
            rows: list = []

        renderer.table(_NoRows())  # nothing to draw, no error
